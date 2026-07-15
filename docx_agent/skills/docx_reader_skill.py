"""DOCX Reader Skill – extracts structured content from a .docx file."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from docx_agent.skills.base_skill import BaseSkill, Issue, SkillResult
from docx_agent.utils.validators import validate_docx_path

# Mapping from WD_ALIGN_PARAGRAPH enum members to readable strings
_ALIGNMENT_MAP: dict[Any, str] = {
    WD_ALIGN_PARAGRAPH.LEFT: "left",
    WD_ALIGN_PARAGRAPH.CENTER: "center",
    WD_ALIGN_PARAGRAPH.RIGHT: "right",
    WD_ALIGN_PARAGRAPH.JUSTIFY: "justify",
    WD_ALIGN_PARAGRAPH.DISTRIBUTE: "distribute",
    None: "inherited",
}


def _rgb_to_hex(color: RGBColor | None) -> str:
    if color is None:
        return "auto"
    return f"#{color.rgb:06X}"


def _pt_value(length: Any) -> float | None:
    """Convert a docx length object to points, or return None."""
    if length is None:
        return None
    try:
        return length.pt
    except AttributeError:
        return None


class DocxReaderSkill(BaseSkill):
    """Extract text, font, and formatting metadata from a .docx document."""

    name = "docx_reader"

    def execute(self, docx_path: str | Path) -> SkillResult:  # type: ignore[override]
        """Read and parse a DOCX file.

        Parameters
        ----------
        docx_path:
            Path to the ``.docx`` file to read.

        Returns
        -------
        SkillResult
            ``metadata`` contains the full structured representation of the
            document under the key ``"document"``.
        """
        path = validate_docx_path(docx_path)
        self._logger.info("Reading DOCX: %s", path)

        doc = Document(str(path))
        issues: list[Issue] = []

        paragraphs_data = self._extract_paragraphs(doc, issues)
        tables_data = self._extract_tables(doc, issues)
        sections_data = self._extract_sections(doc)

        metadata: dict[str, Any] = {
            "file_path": str(path),
            "document": {
                "paragraphs": paragraphs_data,
                "tables": tables_data,
                "sections": sections_data,
                "paragraph_count": len(paragraphs_data),
                "table_count": len(tables_data),
                "word_count": sum(
                    len(p["text"].split()) for p in paragraphs_data if p["text"]
                ),
            },
        }

        self._logger.info(
            "Extracted %d paragraphs, %d tables",
            len(paragraphs_data),
            len(tables_data),
        )
        return SkillResult(
            skill_name=self.name,
            success=True,
            issues=issues,
            metadata=metadata,
        )

    # ------------------------------------------------------------------
    # Private helpers
    # ------------------------------------------------------------------

    def _extract_paragraphs(
        self, doc: Document, issues: list[Issue]
    ) -> list[dict[str, Any]]:
        result = []
        for idx, para in enumerate(doc.paragraphs):
            pf = para.paragraph_format
            para_data: dict[str, Any] = {
                "index": idx,
                "text": para.text,
                "style": para.style.name if para.style else "Normal",
                "alignment": _ALIGNMENT_MAP.get(pf.alignment, "unknown"),
                "space_before_pt": _pt_value(pf.space_before),
                "space_after_pt": _pt_value(pf.space_after),
                "left_indent_pt": _pt_value(pf.left_indent),
                "right_indent_pt": _pt_value(pf.right_indent),
                "first_line_indent_pt": _pt_value(pf.first_line_indent),
                "runs": [],
            }

            for run in para.runs:
                rf = run.font
                run_data: dict[str, Any] = {
                    "text": run.text,
                    "bold": run.bold,
                    "italic": run.italic,
                    "underline": run.underline,
                    "font_name": rf.name,
                    "font_size_pt": _pt_value(rf.size),
                    "font_color": _rgb_to_hex(rf.color.rgb if rf.color and rf.color.type else None),
                    "highlight_color": str(rf.highlight_color) if rf.highlight_color else None,
                }
                para_data["runs"].append(run_data)

            if not para.text.strip() and idx > 0:
                issues.append(
                    Issue(
                        category="format",
                        severity="info",
                        message="Empty paragraph detected.",
                        location=f"Paragraph {idx + 1}",
                    )
                )

            result.append(para_data)

        return result

    def _extract_tables(
        self, doc: Document, issues: list[Issue]
    ) -> list[dict[str, Any]]:
        result = []
        for t_idx, table in enumerate(doc.tables):
            rows_data = []
            for r_idx, row in enumerate(table.rows):
                cells_data = []
                for c_idx, cell in enumerate(row.cells):
                    cells_data.append(
                        {
                            "row": r_idx,
                            "col": c_idx,
                            "text": cell.text,
                            "paragraphs": [p.text for p in cell.paragraphs],
                        }
                    )
                rows_data.append(cells_data)

            if not table.rows:
                issues.append(
                    Issue(
                        category="format",
                        severity="warning",
                        message="Empty table found.",
                        location=f"Table {t_idx + 1}",
                    )
                )

            result.append(
                {
                    "index": t_idx,
                    "row_count": len(table.rows),
                    "col_count": len(table.columns) if table.rows else 0,
                    "rows": rows_data,
                }
            )

        return result

    def _extract_sections(self, doc: Document) -> list[dict[str, Any]]:
        result = []
        for idx, section in enumerate(doc.sections):
            result.append(
                {
                    "index": idx,
                    "page_width_pt": _pt_value(section.page_width),
                    "page_height_pt": _pt_value(section.page_height),
                    "top_margin_pt": _pt_value(section.top_margin),
                    "bottom_margin_pt": _pt_value(section.bottom_margin),
                    "left_margin_pt": _pt_value(section.left_margin),
                    "right_margin_pt": _pt_value(section.right_margin),
                    "orientation": str(section.orientation),
                }
            )
        return result
